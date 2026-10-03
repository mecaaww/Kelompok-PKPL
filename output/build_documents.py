from pathlib import Path
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent
OUT.mkdir(exist_ok=True)

# ID, title, actor, precondition, trigger, main flow, alternative, postcondition
UC = [
('01','Mendaftar akun','Pengunjung','Pengunjung belum masuk ke sistem.','Pengunjung memilih Daftar.',
'1. Pengunjung mengisi nama, email, dan kata sandi. 2. Sistem memeriksa kelengkapan, format dan keunikan email, serta panjang kata sandi. 3. Sistem membuat akun pelanggan, membuat token autentikasi, dan membuka beranda.',
'Email sudah digunakan atau kata sandi kurang dari 8 karakter: sistem menampilkan kesalahan dan meminta perbaikan.',
'Akun pelanggan tersimpan dan pengguna masuk ke sistem.'),
('02','Masuk ke sistem','Pelanggan atau Admin','Akun pengguna sudah terdaftar.','Pengguna membuka halaman Masuk.',
'1. Pengguna memasukkan email dan kata sandi. 2. Sistem memvalidasi kredensial. 3. Sistem menyimpan token pada cookie. 4. Pelanggan diarahkan ke beranda dan admin ke dashboard.',
'Email atau kata sandi salah: sistem menolak autentikasi dan menampilkan pesan kesalahan.',
'Pengguna memiliki sesi autentikasi untuk mengakses fungsi sesuai perannya.'),
('03','Keluar dari sistem','Pelanggan atau Admin','Pengguna memiliki sesi autentikasi.','Pengguna memilih Keluar.',
'1. Pengguna mengirim permintaan keluar. 2. Sistem menjalankan logout dan menghapus cookie token. 3. Sistem mengarahkan pengguna ke beranda.',
'Apabila sesi sudah berakhir, pengguna perlu masuk kembali sebelum memakai fungsi yang dilindungi.',
'Sesi pengguna berakhir dan fungsi pribadi memerlukan autentikasi ulang.'),
('04','Mengatur ulang kata sandi','Pelanggan atau Admin; pendukung Layanan Email','Pengguna memiliki email terdaftar dan dapat mengakses email tersebut.','Pengguna memilih Lupa kata sandi.',
'1. Pengguna memasukkan email. 2. Sistem membuat token reset dan mengirim tautan melalui layanan email. 3. Pengguna membuka tautan dan mengisi kata sandi baru beserta konfirmasinya. 4. Sistem memvalidasi token dan memperbarui kata sandi.',
'Email tidak terdaftar, token tidak valid atau kedaluwarsa, atau konfirmasi kata sandi tidak cocok: sistem menampilkan kesalahan.',
'Kata sandi akun diperbarui jika seluruh validasi berhasil.'),
('05','Menelusuri katalog obat','Pengunjung atau Pelanggan','Data produk tersedia; autentikasi tidak diperlukan.','Aktor membuka beranda atau melakukan pencarian.',
'1. Sistem menampilkan obat yang memiliki stok. 2. Aktor menelusuri halaman, mencari nama atau kategori, atau memakai filter. 3. Sistem menampilkan hasil sesuai kriteria. Beranda menyediakan filter kategori dan awalan nama merek serta daftar produk diskon.',
'Tidak ada produk yang cocok: sistem menampilkan hasil kosong dan aktor dapat mengubah kriteria.',
'Aktor memperoleh daftar obat sesuai kriteria yang dipilih.'),
('06','Melihat detail obat','Pengunjung atau Pelanggan','Produk yang dipilih tersedia dalam data obat.','Aktor memilih sebuah produk.',
'1. Sistem membaca produk dan deskripsi terkait. 2. Sistem menampilkan nama, harga, diskon, gambar, dan informasi obat. 3. Aktor membaca detail sebelum memutuskan pembelian.',
'ID obat tidak ditemukan: sistem mengembalikan halaman tidak ditemukan.',
'Informasi obat ditampilkan tanpa mengubah data produk.'),
('07','Mengelola keranjang','Pelanggan','Pelanggan sudah masuk; produk yang ditambahkan memiliki ID valid.','Pelanggan menambah obat atau membuka keranjang.',
'1. Pelanggan menambahkan obat dengan jumlah minimal satu. 2. Sistem menyimpan item pada keranjang pelanggan. 3. Pelanggan dapat melihat item, menaikkan atau menurunkan jumlah, dan menghapus item. 4. Sistem memperbarui tampilan subtotal.',
'Jumlah turun di bawah satu: item dihapus. Keranjang kosong: tidak ada item untuk dipesan. Item pengguna lain tidak boleh diperbarui.',
'Isi keranjang pelanggan tersimpan sesuai perubahan yang berhasil.'),
('08','Membuat pesanan','Pelanggan','Pelanggan sudah masuk dan keranjang berisi barang.','Pelanggan memilih Lanjut ke Pembayaran.',
'1. Pelanggan menentukan lokasi melalui UC-09 dan melengkapi alamat. 2. Sistem menampilkan ringkasan biaya. 3. Pelanggan mengonfirmasi pesanan. 4. Sistem menyimpan invoice, pesanan, dan detail harga dalam transaksi basis data. 5. Sistem mengosongkan keranjang dan membuka halaman pembayaran.',
'Lokasi atau alamat belum diisi: antarmuka meminta kelengkapan. Keranjang kosong: server menolak pesanan. Pelanggan membatalkan konfirmasi: pesanan tidak dibuat.',
'Pesanan berstatus menunggu pembayaran tersimpan dengan metode Transfer Bank. Relasi: include UC-09.'),
('09','Menentukan lokasi dan ongkir','Pelanggan; pendukung Layanan Peta','Pelanggan sedang menyiapkan pesanan pada halaman keranjang.','Pelanggan membuka pemilih lokasi.',
'1. Sistem memuat peta dan penanda lokasi. 2. Pelanggan menentukan titik pengiriman dan mengonfirmasinya. 3. Sistem menghitung jarak dari apotek, membulatkan ke atas dengan minimum 1 km, lalu mengalikan Rp2.000. 4. Sistem meminta alamat dari koordinat dan memperbarui total.',
'Penerjemahan koordinat ke alamat gagal: kolom alamat diisi koordinat dan dapat dilengkapi pelanggan. Jika izin lokasi tidak diberikan, pelanggan dapat memilih titik secara manual.',
'Koordinat, alamat, ongkir, dan total tersedia untuk UC-08. Jarak merupakan jarak geografis, bukan rute kendaraan.'),
('10','Melihat instruksi pembayaran','Pelanggan','Pelanggan sudah masuk dan memiliki pesanan.','Pelanggan membuka halaman pembayaran setelah checkout.',
'1. Sistem mengambil pesanan milik pelanggan. 2. Sistem menampilkan rincian item, biaya, total bayar, nomor invoice, dan rekening transfer. 3. Pelanggan membaca instruksi untuk melakukan transfer di luar aplikasi.',
'Pesanan tidak ditemukan atau bukan milik pelanggan: akses ditolak melalui pencarian data yang dibatasi pemilik.',
'Instruksi pembayaran ditampilkan. Membuka halaman ini tidak mengubah status pembayaran.'),
('11','Melihat riwayat dan detail pesanan','Pelanggan','Pelanggan sudah masuk.','Pelanggan memilih menu Pesanan.',
'1. Sistem menampilkan pesanan milik pelanggan dari yang terbaru. 2. Pelanggan memilih satu pesanan. 3. Sistem menampilkan invoice, rincian barang, biaya, alamat, dan status pesanan.',
'Belum ada pesanan: daftar kosong ditampilkan. ID pesanan tidak ada atau milik pengguna lain: detail tidak ditampilkan.',
'Pelanggan mengetahui riwayat dan status pesanan tanpa mengubahnya.'),
('12','Melihat dashboard','Admin','Admin sudah masuk ke sistem.','Admin membuka dashboard.',
'1. Sistem menghitung jumlah pesanan, total nilai pesanan selesai, produk dengan stok kurang dari 10, dan akun pelanggan. 2. Sistem menampilkan ringkasan tersebut beserta grafik jumlah pesanan per tanggal.',
'Belum ada transaksi: metrik pesanan bernilai nol dan grafik tidak memiliki data transaksi.',
'Admin memperoleh ringkasan operasional. Jumlah akun pelanggan tidak berarti jumlah pengguna yang sedang online.'),
('13','Mengelola produk obat','Admin','Admin sudah masuk; untuk mengubah atau menghapus, produk harus tersedia.','Admin membuka menu Produk.',
'1. Sistem menampilkan daftar produk dengan pencarian, filter, dan pengurutan. 2. Admin memilih tambah, ubah, atau hapus. 3. Untuk tambah atau ubah, admin mengisi nama, kategori, harga, stok, diskon, dan gambar yang diperlukan. 4. Sistem memvalidasi dan menyimpan perubahan, atau menghapus produk yang dipilih.',
'Data wajib tidak valid: penyimpanan ditolak. Admin membatalkan tindakan: data tetap. Produk tidak ditemukan: operasi tidak dapat dilanjutkan.',
'Data produk diperbarui. Pengelolaan deskripsi obat terstruktur tidak termasuk formulir CRUD ini.'),
('14','Melihat daftar dan detail pesanan','Admin','Admin sudah masuk.','Admin membuka menu Pesanan.',
'1. Sistem menampilkan daftar seluruh pesanan dengan informasi pelanggan. 2. Admin dapat memfilter berdasarkan status. 3. Admin memilih pesanan. 4. Sistem menampilkan pelanggan, item pesanan, alamat, biaya, dan status saat ini.',
'Tidak ada pesanan sesuai filter: daftar kosong. Pesanan tidak ditemukan: detail tidak ditampilkan.',
'Admin memperoleh informasi untuk memproses pesanan.'),
('15','Memperbarui status pesanan','Admin','Admin sudah masuk dan telah memilih pesanan.','Admin memilih status pada detail pesanan.',
'1. Sistem menyediakan status menunggu pembayaran, dikemas, dikirim, selesai, dan dibatalkan. 2. Admin memilih status sesuai hasil penanganan pesanan. 3. Sistem menyimpan status dan menampilkan pesan berhasil.',
'Pesanan tidak ditemukan: pembaruan gagal. Aturan pembatasan perpindahan status merupakan kebutuhan validasi lanjutan.',
'Status terbaru tersimpan dan dapat dilihat pelanggan. Aksi ini bukan verifikasi transfer otomatis.')]

def doc(title, subtitle):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2)
    sec.left_margin=sec.right_margin=Cm(2.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=d.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    s=d.styles['Normal']; s.font.size=Pt(11); s.paragraph_format.space_after=Pt(6); s.paragraph_format.line_spacing=1.08
    for name,size in [('Title',24),('Heading 1',16),('Heading 2',12)]:
        d.styles[name].font.size=Pt(size)
    d.add_paragraph(title,'Title'); d.add_paragraph(subtitle,'Subtitle')
    p=sec.footer.paragraphs[0]; p.alignment=2
    p.add_run('WebApotek  |  ')
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); p._p.append(f)
    d.core_properties.author=''; d.core_properties.title=title
    return d
def p(d,t): d.add_paragraph(t)
def h(d,t): d.add_paragraph(t,'Heading 1')
def sub(d,t): d.add_paragraph(t,'Heading 2')
def page(d,t): d.add_page_break(); h(d,t)
def table(d,heads,rows,widths):
    t=d.add_table(rows=1, cols=len(heads)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i,x in enumerate(heads): t.rows[0].cells[i].text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row): c.text=x
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar); borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');pr.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2)
                for r in pp.runs:r.font.size=Pt(10);r.bold=ri==0
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

d=doc('Spesifikasi Kebutuhan Perangkat Lunak','Bagian a  |  WebApotek Apotek Alfina Rizqy')
p(d,'Nama: ____________________    NIM: ____________________    Kelas: __________')
h(d,'1 Pendahuluan')
p(d,'WebApotek adalah aplikasi pemesanan obat berbasis web untuk Apotek Alfina Rizqy. Aplikasi membantu pelanggan menemukan obat, menyiapkan keranjang, menentukan alamat pengiriman, dan membuat pesanan. Admin menggunakan aplikasi untuk mengelola produk serta memantau dan memperbarui status pesanan.')
p(d,'Dokumen Software Requirements Specification atau SRS ini menetapkan fungsi, data, aturan bisnis, dan kriteria penerimaan sistem sebagai acuan pengembangan dan pengujian. Kebutuhan fungsional diberi ID KF dan dipetakan ke use case UC pada bagian b. Kebutuhan nonfungsional diberi ID KNF dan dinyatakan sebagai target kualitas yang perlu diuji.')
sub(d,'1 1 Ruang lingkup')
p(d,'Ruang lingkup meliputi akun pengguna, katalog dan detail obat, keranjang, lokasi pengiriman dan ongkir, checkout, instruksi transfer bank, riwayat pesanan, dashboard admin, pengelolaan produk, serta perubahan status pesanan.')
p(d,'Pembayaran berlangsung melalui transfer bank di luar aplikasi. Integrasi payment gateway, unggah bukti transfer, konsultasi dokter, unggah resep, pelacakan kurir langsung, dan pengadaan barang dari pemasok tidak termasuk ruang lingkup fitur yang didokumentasikan.')
sub(d,'1 2 Aktor sistem')
table(d,['Aktor','Peran'],[
('Pengunjung','Melihat katalog dan detail obat serta mendaftar akun.'),
('Pelanggan','Mengelola keranjang, membuat pesanan, dan melihat pembayaran serta riwayat pesanan.'),
('Admin','Mengelola produk dan status pesanan serta melihat ringkasan operasional.'),
('Layanan Email','Mengirim tautan untuk pengaturan ulang kata sandi.'),
('Layanan Peta','Menyediakan peta dan penerjemahan koordinat menjadi alamat.')],[3.5,12.9])

page(d,'2 Kebutuhan fungsional akun dan belanja')
reqs=[
('01','Pendaftaran akun','Sistem harus menerima nama, email unik berformat valid, dan kata sandi minimal 8 karakter; membuat akun pelanggan dan mengautentikasi akun setelah pendaftaran berhasil.'),
('02','Autentikasi pengguna','Sistem harus memeriksa email dan kata sandi, memberikan sesi berbasis token, serta mengarahkan pelanggan ke beranda dan admin ke dashboard.'),
('03','Pengakhiran sesi','Sistem harus menyediakan logout yang mengakhiri sesi dan menghapus cookie token pengguna.'),
('04','Pemulihan kata sandi','Sistem harus menyediakan permintaan tautan reset lewat email dan penggantian kata sandi dengan token valid serta konfirmasi kata sandi.'),
('05','Katalog dan pencarian','Sistem harus menampilkan produk dengan stok lebih dari nol, produk diskon, pencarian nama atau kategori, filter kategori, dan navigasi halaman. Beranda juga menyediakan filter merek berdasarkan awalan nama produk.'),
('06','Detail obat','Sistem harus menampilkan informasi produk, harga, diskon, gambar, dan deskripsi obat yang disimpan menurut urutan tampil.'),
('07','Keranjang pelanggan','Sistem harus menyediakan tambah, lihat, perubahan jumlah, dan hapus item keranjang milik pelanggan. Jumlah awal minimal satu; penurunan jumlah di bawah satu menghapus item.'),
('08','Pembuatan pesanan','Sistem harus membuat invoice dan detail pesanan dari keranjang pelanggan, menyimpan alamat serta biaya, menetapkan status awal menunggu pembayaran, dan mengosongkan keranjang setelah penyimpanan berhasil.')]
for id,title,text in reqs:sub(d,f'KF {id} {title}');p(d,text+' Rujukan: UC-'+id+'.')

page(d,'3 Kebutuhan pengiriman dan administrasi')
reqs2=[
('09','Lokasi dan ongkir','Sistem harus menerima titik pengiriman melalui peta, meminta alamat berdasarkan koordinat, serta menghitung ongkir dari jarak geografis ke apotek.'),
('10','Instruksi pembayaran','Sistem harus menampilkan invoice, rincian biaya, metode Transfer Bank, dan informasi rekening bagi pemilik pesanan.'),
('11','Riwayat pelanggan','Sistem harus menampilkan daftar pesanan milik pelanggan dari yang terbaru beserta detail dan statusnya.'),
('12','Dashboard admin','Sistem harus menyajikan total pesanan, nilai pesanan selesai, jumlah obat dengan stok kurang dari 10, jumlah akun pelanggan, dan grafik jumlah pesanan per tanggal.'),
('13','Pengelolaan produk','Sistem harus menyediakan daftar, pencarian, filter, pengurutan, tambah, ubah, dan hapus produk. Data produk meliputi nama, kategori, harga, stok, diskon, dan gambar.'),
('14','Pemeriksaan pesanan oleh admin','Sistem harus menyediakan daftar seluruh pesanan, filter status, dan detail pelanggan, barang, alamat, serta biaya.'),
('15','Pembaruan status pesanan','Sistem harus menyediakan perubahan status pesanan oleh admin dan menampilkan status terkini pada riwayat pelanggan.')]
for id,title,text in reqs2:sub(d,f'KF {id} {title}');p(d,text+' Rujukan: UC-'+id+'.')

page(d,'4 Aturan bisnis dan kebutuhan data')
table(d,['ID','Aturan bisnis'],[
('AB-01','Akun hasil pendaftaran umum memiliki peran pelanggan. Akses administrasi ditujukan hanya bagi admin.'),
('AB-02','Harga setelah diskon = harga awal × (1 − diskon persen / 100). Subtotal item = harga setelah diskon × jumlah.'),
('AB-03','Jarak tagihan adalah jarak dalam km yang dibulatkan ke atas, dengan minimum 1 km. Ongkir = jarak tagihan × Rp2.000. Jarak dihitung dari koordinat apotek ke titik pelanggan.'),
('AB-04','Total bayar = jumlah subtotal seluruh item + ongkir. Checkout membutuhkan keranjang berisi item serta lokasi dan alamat pengiriman.'),
('AB-05','Status awal adalah menunggu pembayaran. Pilihan status admin: menunggu pembayaran, dikemas, dikirim, selesai, dan dibatalkan.'),
('AB-06','Nomor invoice memakai awalan INV, tanggal pembuatan, dan empat karakter acak. Pesanan serta detailnya disimpan dalam transaksi basis data.'),
('AB-07','Pelanggan hanya boleh melihat dan mengubah data pribadi miliknya. Transfer dilakukan di luar aplikasi; tidak ada perubahan status otomatis dari bank.')],[1.7,14.7])
sub(d,'4 1 Data utama')
table(d,['Entitas','Data yang disimpan'],[
('User','Identitas, username, email, kata sandi terhash, dan peran pengguna.'),
('Obat','Identitas, nama, kategori, harga, stok, diskon persen, dan lokasi gambar.'),
('DeskripsiObat','Obat terkait, label informasi, nilai informasi, dan urutan tampil.'),
('Keranjang','Pengguna pemilik, obat terkait, dan jumlah item.'),
('Pesanan','Invoice, pengguna, koordinat, alamat, detail alamat, ongkir, total, status, dan metode pembayaran.'),
('DetailPesanan','Pesanan dan obat terkait, jumlah, harga awal, persentase diskon, harga setelah diskon, dan subtotal.')],[3.5,12.9])
p(d,'Satu pengguna dapat memiliki banyak item keranjang dan banyak pesanan. Satu pesanan memiliki banyak detail pesanan. Setiap detail mengacu pada satu obat; satu obat dapat memiliki banyak informasi deskripsi.')

page(d,'5 Kualitas sistem dan penerimaan')
sub(d,'5 1 Kebutuhan nonfungsional')
p(d,'Ketentuan berikut merupakan target penerimaan, bukan pernyataan bahwa seluruh target telah lulus pengujian.')
table(d,['ID','Target dan cara pemeriksaan'],[
('KNF-01','Keamanan akses: permintaan admin dari akun pelanggan harus ditolak; akses ke pesanan atau keranjang pengguna lain harus ditolak. Uji kedua peran dan data milik pengguna berbeda.'),
('KNF-02','Integritas transaksi: penyimpanan pesanan dan detail harus utuh; kegagalan penyimpanan tidak boleh meninggalkan pesanan parsial. Uji kegagalan transaksi basis data.'),
('KNF-03','Validasi server: jumlah, stok, harga, diskon, ongkir, total, dan status harus diperiksa di server. Uji manipulasi nilai dari browser.'),
('KNF-04','Kegunaan: alur katalog sampai checkout harus dapat digunakan pada lebar layar 360 px dan 1366 px tanpa kontrol utama terpotong. Uji kedua ukuran layar.'),
('KNF-05','Kinerja: target usulan respons katalog dan detail maksimal 3 detik pada 1.000 produk dan 20 pengguna serentak di lingkungan uji yang dicatat.'),
('KNF-06','Ketahanan layanan: kegagalan penerjemahan koordinat ke alamat harus menyediakan isian alamat yang dapat diperbaiki pelanggan.')],[1.9,14.5])
sub(d,'5 2 Skenario penerimaan utama')
p(d,'Uji akun valid dan email ganda (KF-01 sampai KF-04); cari produk dan periksa detailnya (KF-05 sampai KF-06); ubah keranjang lalu buat pesanan (KF-07 sampai KF-10); cocokkan invoice dan status pada riwayat pelanggan (KF-11); tambah atau ubah produk dan perbarui status pesanan melalui admin (KF-12 sampai KF-15).')
p(d,'Contoh perhitungan: dua obat seharga Rp10.000 dengan diskon 10% menghasilkan subtotal Rp18.000. Jarak 2,3 km dibulatkan menjadi 3 km, sehingga ongkir Rp6.000 dan total bayar Rp24.000.')
sub(d,'5 3 Prioritas penyempurnaan implementasi')
p(d,'Pemeriksaan peran admin perlu ditegakkan pada middleware. Perhitungan total dan ongkir perlu divalidasi ulang di server, dan pemeriksaan serta pengurangan stok perlu ditambahkan pada transaksi pemesanan. Validasi perpindahan status dan kesesuaian nama route pengalihan autentikasi juga perlu disempurnakan. Target KNF terkait memerlukan pengujian setelah perbaikan.')
p(d,'Rujukan proyek: routes/web.php; routes/api.php; app/Http/Controllers; app/Http/Middleware/JwtMiddleware.php; app/Models; database/migrations; resources/views/products/keranjang_product.blade.php.')
d.save(OUT/'A_SRS_WebApotek.docx')

b=doc('Deskripsi Use Case WebApotek','Bagian b  |  Apotek Alfina Rizqy')
p(b,'Dokumen ini menjelaskan interaksi aktor dengan WebApotek dan menjadi pasangan diagram dalam file B_Use_Case_WebApotek.drawio. Setiap ID UC pada diagram memiliki deskripsi dengan prasyarat, pemicu, alur utama, kondisi alternatif, dan hasil akhir. ID UC juga berpasangan dengan ID KF pada dokumen SRS.')
h(b,'1 Panduan diagram')
p(b,'Diagram dibagi menjadi tiga halaman agar mudah dibaca dan diedit: Akun dan Katalog, Transaksi Pelanggan, serta Administrasi. Ketiganya merupakan sudut pandang dari sistem WebApotek yang sama. Aktor yang muncul ulang tetap merupakan aktor yang sama.')
p(b,'Garis tanpa panah menunjukkan hubungan aktor dengan use case. Panah putus-putus berlabel «include» dari UC-08 ke UC-09 berarti pembuatan pesanan memerlukan penentuan lokasi dan ongkir. Login dinyatakan sebagai prasyarat fungsi pribadi, bukan langkah yang selalu dijalankan ulang melalui relasi include.')
table(b,['Halaman','Use case'],[
('Akun dan Katalog','UC-01 sampai UC-06: pendaftaran, masuk, keluar, reset kata sandi, katalog, dan detail obat.'),
('Transaksi Pelanggan','UC-07 sampai UC-11: keranjang, checkout, lokasi dan ongkir, instruksi pembayaran, serta riwayat pesanan.'),
('Administrasi','UC-12 sampai UC-15: dashboard, produk, pemeriksaan pesanan, dan pembaruan status.')],[4.4,12])
sub(b,'1 1 Batas sistem dan aktor')
p(b,'Pengunjung, pelanggan, dan admin merupakan aktor manusia. Layanan Email dan Layanan Peta merupakan sistem pendukung di luar batas WebApotek. Pelanggan juga dapat menggunakan katalog dan detail obat yang terbuka bagi pengunjung. Bank dan kurir tidak diberi hubungan sistem karena tidak terdapat integrasi langsung pada alur ini.')
p(b,'Transfer dilakukan di luar aplikasi. Halaman pembayaran hanya menampilkan instruksi, sedangkan admin memperbarui status pesanan secara manual. Hak admin dalam model ini merupakan aturan akses yang harus ditegakkan oleh implementasi.')
sub(b,'1 2 Cara mengedit')
p(b,'Buka diagrams.net atau draw.io, pilih Open Existing Diagram atau File > Open From > Device, kemudian pilih file .drawio. Gunakan tab halaman untuk berpindah diagram. Aktor, oval use case, teks, batas sistem, dan konektor merupakan objek terpisah yang dapat diedit.')
for start in range(0,len(UC),3):
    page(b,f'{2+start//3} Deskripsi UC {start+1:02d} sampai UC {min(start+3,len(UC)):02d}')
    for id,title,actor,pre,trigger,flow,alt,post in UC[start:start+3]:
        sub(b,f'UC {id} {title}')
        for label,value in [('Aktor',actor),('Prasyarat',pre),('Pemicu',trigger),('Alur utama',flow),('Alternatif',alt),('Kondisi akhir',post)]:
            pp=b.add_paragraph();pp.paragraph_format.space_after=Pt(4)
            pp.add_run(label+': ').bold=True;pp.add_run(value)
b.save(OUT/'B_Deskripsi_Use_Case_WebApotek.docx')

# Native editable draw.io XML. Each page is a view of the same system.
mx=ET.Element('mxfile',host='app.diagrams.net',type='device',version='24.7.17')
def diagram(name):
    dg=ET.SubElement(mx,'diagram',id=name.split()[0],name=name)
    model=ET.SubElement(dg,'mxGraphModel',dx='1400',dy='1000',grid='1',gridSize='10',guides='1',tooltips='1',connect='1',arrows='1',fold='1',page='1',pageScale='1',pageWidth='1400',pageHeight='1000',math='0',shadow='0')
    root=ET.SubElement(model,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
    return root
def cell(root,id,value,x,y,w,h,style):
    c=ET.SubElement(root,'mxCell',id=id,value=value,style=style,vertex='1',parent='1')
    ET.SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
def edge(root,id,s,t,label='',include=False,points=None):
    style='edgeStyle=orthogonalEdgeStyle;rounded=0;html=0;endArrow=none;strokeColor=#475569;strokeWidth=1.5;'
    if include:style+='dashed=1;endArrow=open;endFill=0;fontSize=14;labelBackgroundColor=#ffffff;'
    c=ET.SubElement(root,'mxCell',id=id,value=label,style=style,edge='1',parent='1',source=s,target=t)
    g=ET.SubElement(c,'mxGeometry',relative='1',attrib={'as':'geometry'})
    if points:
        a=ET.SubElement(g,'Array',attrib={'as':'points'})
        for x,y in points: ET.SubElement(a,'mxPoint',x=str(x),y=str(y))
def base(root,title,note):
    cell(root,'title',title,30,20,1340,45,'text;html=0;fontSize=26;fontStyle=1;align=center;')
    cell(root,'system','WebApotek Apotek Alfina Rizqy',250,100,850,750,'rounded=0;html=0;fillColor=none;strokeColor=#334155;fontSize=19;verticalAlign=top;spacingTop=16;')
    cell(root,'note',note,250,875,850,80,'text;html=0;whiteSpace=wrap;fontSize=15;align=left;verticalAlign=top;')
def actor(root,id,title,x,y):cell(root,id,title,x,y,90,95,'shape=umlActor;html=0;verticalLabelPosition=bottom;verticalAlign=top;align=center;fontSize=16;fillColor=#ffffff;strokeColor=#334155;')
def use(root,n,x,y):cell(root,'uc'+n,'UC-'+n+'\n'+UC[int(n)-1][1],x,y,280,90,'ellipse;html=0;whiteSpace=wrap;fontSize=16;fillColor=#f4f7fa;strokeColor=#334155;strokeWidth=1.5;')
r=diagram('01 Akun dan Katalog');base(r,'Use Case Akun dan Katalog','Garis menunjukkan asosiasi aktor. Katalog dan detail dapat diakses tanpa login.\nLayanan Email mendukung pengiriman tautan reset kata sandi.')
actor(r,'guest','Pengunjung',65,230);actor(r,'customer','Pelanggan',65,590);actor(r,'admin','Admin',1210,270);actor(r,'email','Layanan Email',1210,645)
for n,x,y in [('01',310,180),('05',310,380),('06',310,590),('02',760,180),('03',760,380),('04',760,590)]:use(r,n,x,y)
for i,(s,t) in enumerate([('guest','01'),('guest','05'),('guest','06'),('customer','05'),('customer','06'),('customer','02'),('customer','03'),('customer','04'),('admin','02'),('admin','03'),('admin','04'),('email','04')]):edge(r,'e'+str(i),s,'uc'+t)
r=diagram('02 Transaksi Pelanggan');base(r,'Use Case Transaksi Pelanggan','Prasyarat: pelanggan sudah masuk. UC-08 menyertakan UC-09.\nTransfer bank berlangsung di luar sistem; UC-10 hanya menampilkan instruksi.')
actor(r,'customer','Pelanggan',65,405);actor(r,'map','Layanan Peta',1210,380)
for n,x,y in [('07',330,180),('08',330,360),('09',770,360),('10',330,540),('11',330,710)]:use(r,n,x,y)
for i,n in enumerate(['07','08','10','11']):edge(r,'e'+str(i),'customer','uc'+n)
edge(r,'incl','uc08','uc09','«include»',True);edge(r,'mapedge','map','uc09')
r=diagram('03 Administrasi');base(r,'Use Case Administrasi','Prasyarat: admin sudah masuk.\nKelola produk mencakup tambah, ubah, hapus, pencarian, filter, dan pengurutan.\nPembaruan status dilakukan secara manual oleh admin.')
actor(r,'admin','Admin',65,400)
for n,x,y in [('12',515,180),('13',515,350),('14',515,520),('15',515,690)]:use(r,n,x,y);edge(r,'e'+n,'admin','uc'+n)
ET.indent(mx,space='  ')
ET.ElementTree(mx).write(OUT/'B_Use_Case_WebApotek.drawio',encoding='utf-8',xml_declaration=True)
for dg in mx:
    cells=dg.findall('.//mxCell');ids=[c.get('id') for c in cells];assert len(ids)==len(set(ids))
    for c in cells:
        if c.get('edge'):assert c.get('source') in ids and c.get('target') in ids
print('Created two DOCX files and one editable draw.io file with 3 pages and 15 use cases.')
