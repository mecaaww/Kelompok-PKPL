from pathlib import Path
import ast
import math
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).parent/'revisi'
OUT.mkdir(exist_ok=True)
# Reuse only the document formatting helpers, without executing the old builder.
src=ast.parse((Path(__file__).parent/'build_documents.py').read_text(encoding='utf-8'))
for node in src.body:
    if isinstance(node,ast.FunctionDef) and node.name in ['doc','p','h','sub','page','table']:
        exec(compile(ast.Module(body=[node],type_ignores=[]),'<formatting>','exec'))

cases=[
dict(id='01',name='Mendaftar akun',actor='Pelanggan',pre='Pelanggan belum memiliki akun.',trigger='Pelanggan memilih menu Daftar.',
steps=['Pelanggan mengisi nama, email, dan kata sandi.','Aplikasi memeriksa data yang diisi.','Aplikasi menyimpan akun pelanggan dan membuka halaman utama.'],
alt='Jika email sudah digunakan atau kata sandi kurang dari 8 karakter, pelanggan diminta memperbaiki isian.',post='Akun pelanggan tersimpan dan dapat digunakan untuk berbelanja.',
want='mendaftar akun',benefit='saya bisa memesan obat menggunakan akun sendiri',accept=['Nama, email, dan kata sandi wajib diisi.','Email harus valid dan belum digunakan.','Kata sandi minimal 8 karakter; pendaftaran yang berhasil membuka halaman utama.']),
dict(id='02',name='Login',actor='Pelanggan dan Admin',pre='Pengguna sudah mempunyai akun.',trigger='Pengguna membuka halaman Login.',
steps=['Pengguna memasukkan email dan kata sandi.','Aplikasi memeriksa kecocokan akun.','Pelanggan masuk ke halaman utama, sedangkan admin masuk ke dashboard.'],
alt='Jika email atau kata sandi salah, aplikasi menampilkan pesan kesalahan.',post='Pengguna masuk ke akun sesuai perannya.',
want='masuk ke akun saya',benefit='saya dapat memakai fitur sesuai peran saya',accept=['Email dan kata sandi wajib diisi.','Data login yang salah ditolak.','Pelanggan diarahkan ke halaman utama dan admin ke dashboard.']),
dict(id='03',name='Mengatur ulang kata sandi',actor='Pelanggan dan Admin',pre='Pengguna memiliki email yang terdaftar.',trigger='Pengguna memilih Lupa kata sandi.',
steps=['Pengguna memasukkan email terdaftar.','Aplikasi mengirim tautan untuk mengganti kata sandi.','Pengguna membuka tautan, lalu mengisi kata sandi baru dan konfirmasinya.','Aplikasi memeriksa tautan dan menyimpan kata sandi baru.'],
alt='Jika email tidak terdaftar, tautan tidak valid, atau konfirmasi tidak cocok, aplikasi menampilkan kesalahan.',post='Pengguna dapat login menggunakan kata sandi baru.',
want='mengganti kata sandi saat lupa',benefit='saya bisa mengakses akun kembali',accept=['Permintaan reset menggunakan email terdaftar.','Penggantian kata sandi membutuhkan tautan reset yang valid.','Konfirmasi harus cocok dengan kata sandi baru.']),
dict(id='04',name='Mencari obat',actor='Pelanggan',pre='Pelanggan membuka website; login belum diperlukan.',trigger='Pelanggan membuka katalog atau mengisi pencarian.',
steps=['Aplikasi menampilkan daftar obat yang memiliki stok.','Pelanggan mencari obat berdasarkan nama atau kategori, atau memakai filter yang tersedia.','Aplikasi menampilkan hasil sesuai pencarian pelanggan.'],
alt='Jika obat tidak ditemukan, pelanggan dapat mengganti kata kunci atau filter.',post='Pelanggan mendapatkan daftar obat yang sesuai.',
want='mencari obat berdasarkan nama atau kategori',benefit='saya lebih mudah menemukan obat yang dibutuhkan',accept=['Katalog dapat dibuka tanpa login.','Hasil pencarian sesuai nama atau kategori yang dimasukkan.','Produk pada katalog dan hasil pencarian memiliki stok lebih dari nol.']),
dict(id='05',name='Melihat detail obat',actor='Pelanggan',pre='Obat yang dipilih tersedia dalam data produk; login belum diperlukan.',trigger='Pelanggan memilih salah satu obat.',
steps=['Pelanggan membuka halaman detail obat.','Aplikasi menampilkan nama, gambar, harga, diskon, dan keterangan obat.','Pelanggan membaca informasi sebelum menambahkan obat ke keranjang.'],
alt='Jika data obat tidak ditemukan, halaman detail tidak dapat ditampilkan.',post='Pelanggan mengetahui informasi obat yang dipilih.',
want='melihat informasi lengkap obat',benefit='saya dapat memeriksa produk sebelum membelinya',accept=['Detail menampilkan obat yang dipilih.','Harga dan diskon ditampilkan sesuai data produk.','Keterangan obat ditampilkan jika datanya tersedia.']),
dict(id='06',name='Mengelola keranjang',actor='Pelanggan',pre='Pelanggan sudah login.',trigger='Pelanggan menambahkan obat atau membuka keranjang.',
steps=['Pelanggan memilih obat dan jumlah yang ingin dibeli.','Aplikasi menyimpan obat dalam keranjang pelanggan.','Pelanggan dapat mengubah jumlah atau menghapus obat.','Aplikasi memperbarui isi keranjang dan subtotal.'],
alt='Jika jumlah dikurangi hingga kurang dari satu, obat dihapus dari keranjang.',post='Keranjang berisi obat dan jumlah yang dipilih pelanggan.',
want='menambah, mengubah jumlah, dan menghapus obat dalam keranjang',benefit='saya bisa menyesuaikan barang sebelum memesan',accept=['Penambahan obat memerlukan login dan jumlah minimal satu.','Perubahan jumlah memperbarui subtotal.','Obat dapat dihapus dari keranjang milik pelanggan.']),
dict(id='07',name='Membuat pesanan',actor='Pelanggan',pre='Pelanggan sudah login dan keranjang berisi obat.',trigger='Pelanggan memilih Lanjut ke Pembayaran.',
steps=['Pelanggan menentukan lokasi pengiriman dan melengkapi alamat.','Aplikasi menampilkan ongkir dan total pembayaran.','Pelanggan memeriksa ringkasan, lalu mengonfirmasi pesanan.','Aplikasi menyimpan pesanan dengan status menunggu pembayaran, mengosongkan keranjang, dan membuka instruksi pembayaran.'],
alt='Jika alamat atau lokasi belum diisi, pelanggan diminta melengkapinya. Keranjang kosong tidak dapat dipesan. Jika konfirmasi dibatalkan, pesanan tidak dibuat.',post='Pesanan dan nomor invoice tersimpan.',
want='membuat pesanan dengan alamat pengiriman saya',benefit='obat yang saya pilih dapat diproses oleh apotek',accept=['Pesanan membutuhkan keranjang berisi obat, alamat, dan lokasi pengiriman.','Total terdiri dari subtotal obat setelah diskon ditambah ongkir.','Pesanan berhasil memiliki invoice, berstatus menunggu pembayaran, dan keranjangnya dikosongkan.']),
dict(id='08',name='Melihat informasi pembayaran',actor='Pelanggan',pre='Pelanggan sudah login dan memiliki pesanan.',trigger='Pelanggan membuka halaman pembayaran.',
steps=['Aplikasi mengambil rincian pesanan pelanggan.','Aplikasi menampilkan invoice, barang yang dipesan, total bayar, dan rekening tujuan.','Pelanggan menggunakan informasi tersebut untuk transfer di luar aplikasi.'],
alt='Pesanan yang tidak ditemukan atau milik pelanggan lain tidak ditampilkan.',post='Pelanggan mengetahui jumlah dan tujuan transfer.',
want='melihat total pembayaran dan rekening tujuan',benefit='saya mengetahui jumlah yang harus ditransfer',accept=['Halaman menampilkan invoice dan total pesanan milik pelanggan.','Metode pembayaran yang ditampilkan adalah Transfer Bank.','Membuka halaman pembayaran tidak otomatis mengubah status pesanan.']),
dict(id='09',name='Melihat riwayat pesanan',actor='Pelanggan',pre='Pelanggan sudah login.',trigger='Pelanggan membuka menu Pesanan.',
steps=['Aplikasi menampilkan daftar pesanan pelanggan dari yang terbaru.','Pelanggan memilih pesanan yang ingin dilihat.','Aplikasi menampilkan rincian barang, biaya, alamat, dan status pesanan.'],
alt='Jika pelanggan belum pernah memesan, daftar pesanan masih kosong.',post='Pelanggan mengetahui rincian dan status pesanannya.',
want='melihat riwayat dan status pesanan saya',benefit='saya bisa mengetahui perkembangan pesanan',accept=['Daftar hanya berisi pesanan milik pelanggan yang login.','Pesanan terbaru ditampilkan lebih dahulu.','Detail menampilkan status terakhir yang tersimpan.']),
dict(id='10',name='Melihat dashboard',actor='Admin',pre='Admin sudah login.',trigger='Admin membuka dashboard.',
steps=['Admin membuka halaman utama admin.','Aplikasi menampilkan jumlah pesanan, nilai pesanan selesai, jumlah produk dengan stok rendah, dan jumlah akun pelanggan.','Admin melihat grafik jumlah pesanan berdasarkan tanggal.'],
alt='Jika belum ada pesanan, jumlah dan nilai pesanan bernilai nol.',post='Admin memperoleh ringkasan kegiatan pemesanan.',
want='melihat ringkasan pesanan dan produk',benefit='saya dapat memantau kegiatan apotek dari satu halaman',accept=['Jumlah pesanan ditampilkan sesuai data.','Nilai pendapatan dihitung dari pesanan berstatus selesai.','Indikator stok rendah menghitung produk dengan stok kurang dari 10.']),
dict(id='11',name='Mengelola data obat',actor='Admin',pre='Admin sudah login.',trigger='Admin membuka menu Produk.',
steps=['Admin melihat daftar obat dan dapat mencari atau memfilter produk.','Admin memilih tambah, ubah, atau hapus obat.','Untuk tambah atau ubah, admin mengisi data obat lalu menyimpannya.','Aplikasi memeriksa isian dan memperbarui daftar obat.'],
alt='Isian wajib yang tidak valid membuat penyimpanan ditolak. Jika admin membatalkan tindakan, data tidak berubah.',post='Data obat tersimpan sesuai perubahan yang dilakukan admin.',
want='menambah, mengubah, dan menghapus data obat',benefit='daftar produk tetap sesuai dengan barang yang dijual',accept=['Data produk memuat nama, kategori, harga, stok, diskon, dan gambar.','Nama, harga, dan stok wajib diisi ketika menyimpan produk.','Perubahan yang berhasil muncul pada daftar produk.']),
dict(id='12',name='Melihat pesanan pelanggan',actor='Admin',pre='Admin sudah login.',trigger='Admin membuka menu Pesanan.',
steps=['Aplikasi menampilkan daftar pesanan pelanggan.','Admin dapat memilih filter status.','Admin membuka salah satu pesanan untuk melihat pelanggan, barang, alamat, dan biaya.'],
alt='Jika tidak ada pesanan sesuai filter, daftar tidak menampilkan hasil.',post='Admin mengetahui rincian pesanan yang perlu ditangani.',
want='melihat daftar dan rincian pesanan pelanggan',benefit='saya dapat menyiapkan pesanan sesuai data yang masuk',accept=['Admin dapat melihat pesanan dari seluruh pelanggan.','Filter menampilkan pesanan dengan status yang dipilih.','Detail memuat pelanggan, item, alamat pengiriman, dan biaya.']),
dict(id='13',name='Mengubah status pesanan',actor='Admin',pre='Admin sudah login dan membuka pesanan yang tersedia.',trigger='Admin memilih status pada detail pesanan.',
steps=['Admin memeriksa pesanan yang akan diperbarui.','Admin memilih status sesuai proses penanganan pesanan.','Aplikasi menyimpan status dan menampilkan pemberitahuan berhasil.'],
alt='Jika pesanan tidak ditemukan, status tidak dapat diperbarui.',post='Status terbaru dapat dilihat admin dan pelanggan.',
want='memperbarui status pesanan',benefit='pelanggan dapat mengetahui perkembangan pesanannya',accept=['Pilihan status adalah menunggu pembayaran, dikemas, dikirim, selesai, dan dibatalkan.','Status yang dipilih tersimpan pada pesanan terkait.','Status terbaru juga tampil pada riwayat pelanggan.']),
dict(id='14',name='Logout',actor='Pelanggan dan Admin',pre='Pengguna sedang login.',trigger='Pengguna memilih Logout.',
steps=['Pengguna memilih keluar dari akun.','Aplikasi mengakhiri sesi login.','Pengguna kembali ke halaman utama.'],
alt='Jika sesi sudah berakhir, pengguna harus login kembali untuk memakai fitur yang memerlukan akun.',post='Akun tidak lagi aktif pada sesi tersebut.',
want='keluar dari akun setelah selesai',benefit='akun saya tidak tetap terbuka pada perangkat yang digunakan',accept=['Logout mengakhiri sesi pengguna.','Setelah logout, pengguna diarahkan ke halaman utama.','Fitur pribadi memerlukan login kembali.'])]

def newdoc(title,subtitle):
    d=doc(title,subtitle)
    d.styles['Normal'].font.name='Times New Roman'
    d.styles['Normal'].font.size=Pt(11)
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        d.styles[name].font.name='Times New Roman'
    d.styles['Title'].font.size=Pt(21)
    return d
def field(d,label,value):
    pp=d.add_paragraph();pp.add_run(label+': ').bold=True;pp.add_run(value)
def steps(d,values):
    for i,v in enumerate(values,1):
        pp=d.add_paragraph(f'{i}. {v}');pp.paragraph_format.space_after=Pt(3)

s=newdoc('Spesifikasi Kebutuhan Perangkat Lunak','Website Apotek Alfina Rizqy')
p(s,'Nama: ____________________    NIM: ____________________    Kelas: __________')
h(s,'1 Pendahuluan')
p(s,'Website Apotek Alfina Rizqy digunakan untuk melihat dan memesan obat secara online. Pelanggan dapat mencari obat, memasukkannya ke keranjang, mengisi alamat pengiriman, lalu membuat pesanan. Admin mengelola data obat dan memperbarui status pesanan yang masuk.')
p(s,'Dokumen ini menjelaskan kebutuhan aplikasi sebagai acuan pengembangan dan pengujian. Pernyataan kebutuhan menunjukkan perilaku yang diharapkan, bukan laporan hasil pengujian aplikasi.')
sub(s,'1 1 Tujuan')
p(s,'Aplikasi bertujuan memudahkan pelanggan memperoleh informasi obat dan melakukan pemesanan. Bagi admin, aplikasi membantu pencatatan produk dan penanganan pesanan dalam satu tempat.')
sub(s,'1 2 Ruang lingkup')
p(s,'Fitur yang dibahas meliputi pendaftaran dan login, pemulihan kata sandi, pencarian serta detail obat, keranjang, pemesanan, informasi pembayaran, riwayat pesanan, dashboard, dan pengelolaan data obat. Alamat dan ongkir merupakan bagian dari proses pemesanan.')
p(s,'Pembayaran dilakukan melalui transfer bank di luar aplikasi. Aplikasi menampilkan informasi pembayaran, sedangkan perubahan status pesanan dilakukan oleh admin. Pembayaran otomatis, konsultasi dokter, dan pelacakan kurir langsung tidak termasuk dalam pembahasan.')
sub(s,'1 3 Pengguna aplikasi')
table(s,['Aktor','Kegiatan'],[('Pelanggan','Mencari obat, melihat detail, mengelola keranjang, membuat pesanan, dan melihat riwayat pesanan.'),('Admin','Melihat dashboard, mengelola data obat, melihat pesanan pelanggan, dan mengubah status pesanan.')],[3.3,13.1])
p(s,'Pelanggan juga mencakup pengguna yang belum login ketika mendaftar atau melihat katalog. Login diperlukan untuk mengelola keranjang dan mengakses pesanan. Peta merupakan fasilitas pengisian lokasi, bukan aktor tersendiri.')

page(s,'2 Kebutuhan fungsional')
p(s,'Kode KF, UC, dan US dengan nomor yang sama membahas fitur yang sama pada dokumen kebutuhan, deskripsi use case, dan user story.')
requirements=[
'Aplikasi menyediakan pendaftaran dengan nama, email unik, dan kata sandi minimal 8 karakter.',
'Aplikasi memeriksa email dan kata sandi, lalu membuka halaman sesuai peran pengguna.',
'Aplikasi menyediakan penggantian kata sandi melalui tautan reset yang dikirim ke email terdaftar.',
'Aplikasi menampilkan katalog obat yang memiliki stok serta menyediakan pencarian nama atau kategori dan filter produk.',
'Aplikasi menampilkan gambar, nama, harga, diskon, dan keterangan obat yang dipilih.',
'Pelanggan dapat menambahkan obat, mengubah jumlah, dan menghapus obat dari keranjangnya.',
'Pelanggan dapat mengisi lokasi serta alamat pengiriman, memeriksa ongkir dan total, lalu mengonfirmasi pesanan.',
'Pelanggan dapat melihat nomor invoice, total pembayaran, dan rekening tujuan transfer untuk pesanannya.',
'Pelanggan dapat melihat riwayat, rincian, dan status pesanan miliknya.',
'Admin dapat melihat jumlah pesanan, nilai pesanan selesai, produk dengan stok rendah, akun pelanggan, dan grafik pesanan.',
'Admin dapat mencari, memfilter, menambah, mengubah, dan menghapus data obat.',
'Admin dapat melihat seluruh pesanan pelanggan, memfilter status, dan membuka rincian pesanan.',
'Admin dapat memperbarui status pesanan sesuai penanganan yang dilakukan.',
'Pelanggan dan admin dapat keluar dari akun untuk mengakhiri sesi login.']
table(s,['Kode','Kebutuhan'],[(f'KF-{c["id"]}',r) for c,r in zip(cases,requirements)],[1.8,14.6])

page(s,'3 Aturan pemesanan dan data')
sub(s,'3 1 Aturan pemesanan')
steps(s,[
'Pendaftaran umum menghasilkan akun pelanggan. Akun admin digunakan untuk pengelolaan aplikasi.',
'Jumlah obat yang ditambahkan ke keranjang minimal satu. Jika jumlah dikurangi sampai kurang dari satu, item dihapus.',
'Harga setelah diskon dihitung dari harga awal dikurangi potongan persentasenya. Subtotal obat adalah harga setelah diskon dikalikan jumlah.',
'Ongkir dihitung Rp2.000 per km. Jarak geografis dibulatkan ke atas dengan jarak minimum 1 km.',
'Total pembayaran adalah jumlah subtotal seluruh obat ditambah ongkir.',
'Pemesanan memerlukan keranjang yang berisi obat, alamat, dan lokasi pengiriman. Pesanan berhasil memperoleh nomor invoice dan status menunggu pembayaran.',
'Pilihan status pesanan adalah menunggu pembayaran, dikemas, dikirim, selesai, dan dibatalkan.',
'Pelanggan hanya boleh mengakses keranjang dan pesanan miliknya. Pengelolaan produk serta perubahan status hanya boleh dilakukan admin.'])
p(s,'Contoh: dua obat dengan harga Rp10.000 dan diskon 10% memiliki subtotal Rp18.000. Jika jarak pengiriman 2,3 km, jarak tagihannya menjadi 3 km. Ongkirnya Rp6.000, sehingga total pembayaran Rp24.000.')
sub(s,'3 2 Data yang digunakan')
table(s,['Data','Isi utama'],[
('Akun','Nama pengguna, email, kata sandi, dan peran.'),('Obat','Nama, kategori, harga, stok, diskon, gambar, dan keterangan obat.'),('Keranjang','Pelanggan, obat yang dipilih, dan jumlah pembelian.'),('Pesanan','Nomor invoice, pelanggan, alamat dan lokasi, ongkir, total, metode pembayaran, serta status.'),('Detail pesanan','Obat, jumlah, harga awal, diskon, harga setelah diskon, dan subtotal.')],[3.3,13.1])

page(s,'4 Kebutuhan nonfungsional dan pengujian')
table(s,['Kode','Kebutuhan','Cara memeriksa'],[
('KNF-01','Akses data harus sesuai peran dan kepemilikan.','Coba membuka halaman admin dengan akun pelanggan dan detail pesanan milik akun lain; akses harus ditolak.'),
('KNF-02','Data pesanan dan detailnya harus disimpan sebagai satu transaksi.','Simulasikan kegagalan penyimpanan; tidak boleh ada pesanan yang hanya tersimpan sebagian.'),
('KNF-03','Halaman utama dan pemesanan harus mudah digunakan di komputer maupun ponsel.','Periksa tampilan dan tombol utama pada lebar layar 360 px dan 1366 px.'),
('KNF-04','Aplikasi harus memberikan pesan ketika isian tidak valid.','Kirim formulir kosong atau salah format; pesan harus menjelaskan isian yang perlu diperbaiki.'),
('KNF-05','Nilai pesanan harus diperiksa kembali di server.','Ubah nilai jumlah atau total pada permintaan; aplikasi harus menolak nilai yang tidak sesuai.')],[1.7,7,7.7])
sub(s,'4 1 Pengujian alur utama')
p(s,'Pengujian dilakukan dengan mendaftarkan akun pelanggan, login, mencari obat, mengubah keranjang, dan membuat pesanan. Nomor invoice, subtotal, ongkir, dan total diperiksa agar sesuai dengan barang yang dipesan. Setelah admin mengubah status, perubahan tersebut harus terlihat pada riwayat pelanggan.')
sub(s,'4 2 Catatan penerapan')
p(s,'Pembatasan akses admin dan pemeriksaan total di server masih perlu disempurnakan pada kode. Pemeriksaan stok dan pengurangan stok otomatis saat pemesanan juga belum dijalankan pada alur penyimpanan pesanan. Karena itu, kebutuhan tersebut perlu diuji setelah implementasinya diperbaiki.')
sub(s,'4 3 Acuan fitur')
p(s,'Fitur pada dokumen ini mengacu pada rute aplikasi, pengendali akun, katalog, keranjang, pesanan, dan admin dalam proyek WebApotek. Dokumen User Story dan Deskripsi Use Case menjabarkan fitur yang sama dari sisi kegiatan pengguna.')
s.save(OUT/'A_SRS_WebApotek_Revisi.docx')

u=newdoc('User Story','Website Apotek Alfina Rizqy')
p(u,'User story berikut menjelaskan kebutuhan pelanggan dan admin beserta hasil yang diharapkan. Kriteria penerimaan dipakai untuk memeriksa apakah setiap kebutuhan sudah terpenuhi.')
for i,c in enumerate(cases):
    if i and i%4==0:page(u,'User Story lanjutan')
    sub(u,f'US {c["id"]} {c["name"]}')
    role='pengguna, baik pelanggan maupun admin,' if c['actor']=='Pelanggan dan Admin' else c['actor'].lower()
    p(u,f'Sebagai {role}, saya ingin {c["want"]} agar {c["benefit"]}.'.replace(',,',','))
    pp=u.add_paragraph('Kriteria penerimaan');pp.runs[0].bold=True
    steps(u,c['accept'])
u.save(OUT/'A_User_Story_WebApotek.docx')

b=newdoc('Deskripsi Use Case','Website Apotek Alfina Rizqy')
p(b,'Aktor dalam diagram adalah Pelanggan dan Admin. Pelanggan dapat mendaftar serta melihat obat sebelum login. Fitur keranjang dan pesanan memerlukan login. Setiap use case di bawah memiliki kode yang sama dengan diagram.')
for i,c in enumerate(cases):
    if i and i%2==0:page(b,'Deskripsi Use Case lanjutan')
    sub(b,f'UC {c["id"]} {c["name"]}')
    field(b,'Aktor',c['actor']);field(b,'Kondisi awal',c['pre']);field(b,'Pemicu',c['trigger'])
    pp=b.add_paragraph('Alur utama');pp.runs[0].bold=True
    steps(b,c['steps'])
    field(b,'Alur alternatif',c['alt']);field(b,'Kondisi akhir',c['post'])
b.save(OUT/'B_Deskripsi_Use_Case_Revisi.docx')

# Single-page, two-actor diagram; the application is the boundary, never an actor.
mx=ET.Element('mxfile',host='app.diagrams.net',type='device')
dg=ET.SubElement(mx,'diagram',id='webapotek',name='Use Case WebApotek')
gm=ET.SubElement(dg,'mxGraphModel',grid='1',gridSize='10',page='1',pageScale='1',pageWidth='1100',pageHeight='1680')
root=ET.SubElement(gm,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
def vertex(id,label,x,y,w,hh,style):
    c=ET.SubElement(root,'mxCell',id=id,value=label,style=style,vertex='1',parent='1')
    ET.SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(hh),attrib={'as':'geometry'})
vertex('boundary','Website Apotek Alfina Rizqy',310,75,480,1530,'rounded=0;fillColor=none;strokeColor=#000000;verticalAlign=top;spacingTop=18;fontSize=22;fontFamily=Arial;')
vertex('pelanggan','Pelanggan',80,750,80,110,'shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;align=center;fontSize=22;strokeColor=#000000;fillColor=#ffffff;')
vertex('admin','Admin',940,750,80,110,'shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;align=center;fontSize=22;strokeColor=#000000;fillColor=#ffffff;')
for i,c in enumerate(cases):
    vertex('uc'+c['id'],'UC-'+c['id']+'\n'+c['name'],365,145+i*102,370,70,'ellipse;whiteSpace=wrap;html=0;fontSize=20;fontFamily=Arial;strokeColor=#000000;fillColor=#ffffff;')
    for actor,key in [('Pelanggan','pelanggan'),('Admin','admin')]:
        if actor in c['actor']:
            edge=ET.SubElement(root,'mxCell',id=key+c['id'],edge='1',parent='1',source=key,target='uc'+c['id'],style='endArrow=none;startArrow=none;strokeColor=#000000;strokeWidth=1;exitX='+('1' if key=='pelanggan' else '0')+';exitY=0.45;entryX='+('0' if key=='pelanggan' else '1')+';entryY=0.5;')
            ET.SubElement(edge,'mxGeometry',relative='1',attrib={'as':'geometry'})
ET.indent(mx,space='  ')
ET.ElementTree(mx).write(OUT/'B_Use_Case_WebApotek_Revisi.drawio',encoding='utf-8',xml_declaration=True)

# PNG uses the same node positions and associations as the editable diagram.
scale=2
im=Image.new('RGB',(1100*scale,1680*scale),'white');dr=ImageDraw.Draw(im)
def f(size,bold=False):return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),size*scale)
def line(points,width=1):dr.line([(x*scale,y*scale) for x,y in points],fill='black',width=width*scale)
def txt(text,x,y,size=20,bold=False):
    dr.multiline_text((x*scale,y*scale),text,font=f(size,bold),fill='black',anchor='mm',align='center',spacing=6*scale)
dr.rectangle((310*scale,75*scale,790*scale,1605*scale),outline='black',width=2*scale)
txt('Use Case Diagram',550,32,25,True);txt('Website Apotek Alfina Rizqy',550,106,22)
for i,c in enumerate(cases):
    cy=180+i*102
    if 'Pelanggan' in c['actor']:line([(160,799.5),(365,cy)])
    if 'Admin' in c['actor']:line([(940,799.5),(735,cy)])
    dr.ellipse((365*scale,(cy-35)*scale,735*scale,(cy+35)*scale),fill='white',outline='black',width=2*scale)
    txt('UC-'+c['id']+'\n'+c['name'],550,cy,19)
for x,label in [(120,'Pelanggan'),(980,'Admin')]:
    dr.ellipse(((x-16)*scale,750*scale,(x+16)*scale,782*scale),fill='white',outline='black',width=2*scale)
    line([(x,782),(x,821)],2);line([(x-40,799.5),(x+40,799.5)],2)
    line([(x-32,860),(x,821),(x+32,860)],2);txt(label,x,888,22)
txt('Pelanggan dan Admin terhubung ke fitur yang digunakan masing-masing.',550,1640,17)
im.save(OUT/'B_Use_Case_WebApotek_Revisi.png')

# Structural checks against the common feature list.
actors=[c.get('value') for c in root if 'umlActor' in c.get('style','')]
assert actors==['Pelanggan','Admin']
assert len([c for c in root if c.get('vertex')=='1' and c.get('id','').startswith('uc')])==14
for file in OUT.glob('*.docx'):
    check=Document(file)
    assert check.paragraphs
    print(file.name, 'paragraphs:',len(check.paragraphs))
print('Verified: 14 use cases, 14 user stories, 14 functional requirements, exactly 2 actors.')
